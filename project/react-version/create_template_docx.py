#!/usr/bin/env python3
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def build_template_docx(output_path):
    doc = Document()

    # Document Title & Subtitle
    title_p = doc.add_paragraph("MANUAL BOOK PANDUAN PENGGUNAAN SISTEM AUTOMASI AI")
    title_p.style.font.size = Pt(22)
    title_p.style.font.bold = True
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    sub_p = doc.add_paragraph("Dokumentasi Sistematis Arsitektur, Fitur, dan Panduan Operasional Perangkat")
    sub_p.style.font.size = Pt(13)
    sub_p.style.font.italic = True

    doc.add_paragraph() # Spacing

    # Section 1
    h1_1 = doc.add_heading("1. Pengenalan Umum Sistem", level=1)
    
    h2_1 = doc.add_heading("1.1 Overview Aplikasi", level=2)
    doc.add_paragraph("Bab ini memuat gambaran umum, tata cara konfigurasi, serta arsitektur alur kerja otomasi berbasis kecerdasan buatan untuk pemrosesan dokumen otomatis.")

    h2_2 = doc.add_heading("1.2 Spesifikasi Parameter Modul", level=2)
    doc.add_paragraph("MaxConcurrency : Jumlah maksimum tugas paralel yang dapat diproses secara bersamaan oleh worker node.")
    doc.add_paragraph("TimeoutPeriod : Batas waktu maksimal (dalam detik) sebelum sistem menandakan status kegagalan koneksi.")
    doc.add_paragraph("RetryAttempts : Batas percobaan ulang eksekusi ketika terjadi hambatan jaringan sementara.")

    doc.add_paragraph("Fungsi Transmit Power : Mengatur kekuatan daya pancar sinyal radio modul jika dinaikkan jangkauan sinyal meluas jika diturunkan menghemat konsumsi daya.")

    doc.add_paragraph("• Mendukung integrasi webhook secara langsung.")
    doc.add_paragraph("• Ekstraksi teks berbasis OCR presisi tinggi.")
    doc.add_paragraph("- Penyimpanan log otomatis pada server lokal.")

    doc.add_paragraph("Catatan: Pastikan lisensi perangkat masih berlaku sebelum melakukan pembaruan firmware.")

    doc.add_paragraph() # Spacing

    # Section 2
    h1_2 = doc.add_heading("2. Panduan Penggunaan dan Prosedur Lapangan", level=1)

    h2_3 = doc.add_heading("2.1 Tahapan Instalasi dan Konfigurasi Node", level=2)
    doc.add_paragraph("Ikuti urutan langkah berikut untuk mengaktifkan modul otomatisasi:")
    doc.add_paragraph("1. Unduh pustaka ketergantungan utama menggunakan package manager.")
    doc.add_paragraph("2. Inisialisasi berkas konfigurasi lingkungan (.env) sesuai dengan kunci API yang valid.")
    doc.add_paragraph("3. Jalankan perintah migrasi skema database lokal untuk penyiapan awal.")

    h2_4 = doc.add_heading("2.2 Rincian Alur Kerja Per Eksekusi", level=2)
    doc.add_paragraph("Langkah 1 : Validasi otentikasi kunci enkripsi pada header permintaan.")
    doc.add_paragraph("Langkah 2 : Pengambilan data mentah dari antrean pesan (message queue).")
    doc.add_paragraph("Langkah 3 : Transformasi data ke dalam format JSON terstruktur.")

    doc.add_paragraph() # Spacing

    # Section 3
    h1_3 = doc.add_heading("3. Spesifikasi Teknis dan Grafik Data", level=1)

    h2_5 = doc.add_heading("3.1 Tabel Pemetaan Kinerja Sistem", level=2)

    # Add Table
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Modul Utama'
    hdr_cells[1].text = 'Kapasitas Beban'
    hdr_cells[2].text = 'Status Operasional'

    data = [
        ('Engine Ingest', '1.500 req/sec', 'Stabil'),
        ('Parser PDF', '850 req/sec', 'Optimal'),
        ('Exporter PDF', '1.200 req/sec', 'Aktif'),
    ]
    for i, row_data in enumerate(data):
        row_cells = table.rows[i+1].cells
        row_cells[0].text = row_data[0]
        row_cells[1].text = row_data[1]
        row_cells[2].text = row_data[2]

    doc.add_paragraph() # Spacing

    doc.add_paragraph("Gambar 1: Caption Grafik Respon Transmisi Data Sistem Automasi")

    doc.add_paragraph() # Spacing

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"✅ Generated template docx at {output_path}")

if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'public', 'Template_Format_ManualBook.docx')
    build_template_docx(out)
